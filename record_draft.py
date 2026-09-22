import json
import requests
from time import sleep
from pathlib import Path

import openpyxl
import pandas as pd
from openpyxl import styles
from itertools import chain

import helper
import constants

json_input_path = Path.cwd().joinpath("input", "json", "draft_input.json")
xlsx_path = Path.cwd().joinpath("input", "xlsx")
output_path = Path.cwd().joinpath("output")
text_path = Path.cwd().joinpath("input", "text")


def record_xlsx_draft(format="json") -> None:
    """"""
    if format == "json":
        with json_input_path.open('r', encoding="utf-8") as f:
            draft_input = json.loads(f.read())
    else:
        draft_input = text_to_dict()

    #### spell_check(draft_input)

    wb = openpyxl.Workbook()
    base_style = styles.NamedStyle(name="base", font=styles.Font(name="Cambria", size=12))
    # Creates a SEAT_ORDER based on which seats were used in the draft
    color_circle = sorted([c for c in draft_input["color_map"]],
                          key=constants.SEAT_ORDER.index)

    def write_cell(sheet, row, col, value, color = None, justify = None):
        """"""
        curr_cell = sheet.cell(row=row, column=col)
        curr_cell.value = value
        curr_cell.style = base_style
        curr_cell.alignment = styles.Alignment(horizontal="center" if not justify else justify)
        if color:
            curr_cell.fill = styles.PatternFill(start_color=constants.COLOR_HEX[color], fill_type="solid")

    # DRAFT SHEET
    draft_sheet = wb.active
    draft_sheet.title = "Draft"
    for player, pick_order in draft_input["player_order"].items():
        # Reverse search of dictionary for player : color
        color = next((key for key, val in draft_input["color_map"].items()
                      if val == player))
        column = color_circle.index(color) + 1

        write_cell(draft_sheet, 1, column, player, color)
        for row, card in enumerate(pick_order, start=0):
            row_offset = row // (draft_input["pack_size"]) # Row break for each pack
            write_cell(draft_sheet, row + row_offset + 2, column, card, color) # 2, 1 for header, 1 for xlsx is 1-indexed

    for color, pick_order in draft_input["color_order"].items():
        player = draft_input["color_map"][color]
        column = color_circle.index(color) + draft_input["num_players"] + 2
        curr_col_idx = color_circle.index(color)
        direction = 1 # Used to represent passing direction

        write_cell(draft_sheet, 1, column, player, color)
        for row, card in enumerate(pick_order, start=0):
            row_offset = row // (draft_input["pack_size"])
            write_cell(draft_sheet, 2 + row + row_offset, column, card,
                       color_circle[curr_col_idx])

            if (row + 1) % draft_input["pack_size"] == 0: # Alternates coloring when the passing direction changes
                curr_col_idx = color_circle.index(color)
                direction = direction * -1
            else:
                curr_col_idx = (curr_col_idx + direction) % len(color_circle)


    # Sort Decklists on cmc
    oracle_data = helper.get_oracle_data()
    deck_sheet = wb.create_sheet(title="Play")
    for player, decklist in draft_input["decklists"].items():
        draft_input["decklists"][player] = sorted(decklist,
                                                 key=lambda x: (oracle_data[x]["cmc"], x))
        color = next((key for key, val in draft_input["color_map"].items()
                      if val == player))
        column = color_circle.index(color) + 1
        write_cell(deck_sheet, 1, column, player, color)
        for row, card in enumerate(draft_input["decklists"][player], start=2):
            write_cell(deck_sheet, row, column, card, color)
        
        companion_column = column + draft_input["num_players"] + 1
        write_cell(deck_sheet, 1, companion_column, player, color)
        if player in draft_input["companions"]:
            for row, card in enumerate(draft_input["companions"][player], start=2):
                 write_cell(deck_sheet, row, companion_column, card, color)

    # RESULTS SHEET
    results_sheet = wb.create_sheet(title="Results")
    write_cell(results_sheet, 1, 1, "Player 1")
    write_cell(results_sheet, 1, 3, "Player 2")

    for entry, record in draft_input["results"].items():
        match, player = entry.strip("m").split("p")
        col = 3 if int(player) == 2 else 1
        color = next((key for key, val in draft_input["color_map"].items()
                      if val == record["player"]))
        write_cell(results_sheet, int(match) + 1, col, record["player"], color)
        write_cell(results_sheet, int(match) + 1, col + 1, record["win_count"], color)

    # META SHEET
    meta_sheet = wb.create_sheet(title="Meta")
    header = {"Draft Number" :"draft",
              "Date" : "date",
              "Patch" : "patch",
              "Pack Size" : "pack_size",
              "Number of Players" : "num_players",
              "Rounds" : "rounds",
              "Draft Type" : "draft_type",
              "Notes" : "notes"}
    
    for row, (key, value) in enumerate(header.items(), start=1):
        write_cell(meta_sheet, row, 1, key, justify="left")
        write_cell(meta_sheet, row, 2, draft_input[value], justify="right")

    

    

    wb.save(output_path.joinpath("wartata.xlsx"))
    print()
        

# spellchecks the cards in color_order of draft_input
def spell_check(draft_input, check="local"):

    combined = [x for player in draft_input["color_order"].values() for x in player]
    if check == "local":
        oracle_data = helper.get_oracle_data()
        for card in combined:
            if card not in oracle_data:
                print(card)
    elif check == "requests":
        confirm = int(input("Check for misspellings? 1 - All, 2 - Ignore Short, 3 - Skip: "))
        if confirm != 3:
            for card in combined:
                request_string = f"https://api.scryfall.com/cards/search?q={card}"
                response = requests.get(request_string,
                                        params = {"format": "json"},
                                        headers = {"User-Agent": constants.USER_AGENT})
                try:
                    card_dict = response.json()["data"]
                    for card_entry in card_dict:
                        if card != card_entry['name'] and confirm == 1:
                            print(f"{card} -- {card_entry['name']}")
                except:
                    print(card)
                sleep(constants.SCRYFALL_SEARCH_RATE)
                

def text_to_dict() -> dict:
    """"""
    with (text_path.joinpath("player_input.txt").open(encoding="utf-8") as f1,
          text_path.joinpath("pack_input.txt").open(encoding="utf-8") as f2,
          text_path.joinpath("deck_input.txt").open(encoding="utf-8") as f3):
        player_input = f1.read().split('\n')
        pack_input = f2.read().split('\n')
        deck_input = f3.read().split('\n')

    # Recording player picks and color to player mapping
    player_order : dict[str, dict] = {}
    color_map : dict[str, str] = {}
    for line in player_input:
        if not line:
            continue
        elif "-#-" in line:
            curr_player, curr_color = line.split("-#-")
            color_map[curr_color] = curr_player
            player_order[curr_player] = []
        else:
            player_order[curr_player].append(line)
    # Combines each player order while maintaining pick order,
    # Then maps that pick order to its subsequent index 0-179 
    combined = {card: idx for idx, card in
                enumerate(chain.from_iterable(zip(*player_order.values())))}
    packs: list[list] = [[]]
    sorted_packs : list[list] = []
    # If there are many empty strings in a row, will skip all but the first.
    # This is to account for occasional multiple empty strings at end of file.
    empty_line = False
    for line in pack_input:
        if not line:
            if empty_line:
                continue
            else:
                # Sorts each pack independently based on combined
                sorted_packs.append(sorted(packs[len(packs) - 1], 
                                        key = lambda x: combined[x]))
                packs.append([])
                empty_line = True
        else:
            packs[len(packs) - 1].append(line)
            empty_line = False
    color_order: dict[str, list] = {key: [] for key in color_map.keys()}
    for pack in sorted_packs:
        # find color from first pick
        fp = pack[0]
        for color, player in color_map.items():
            if fp in player_order[player]:
                color_order[color].extend(pack)
    decklists: dict[str, list] = {}
    companions: dict[str, list] = {}
    has_companions = False
    for line in deck_input:
        if not line:
            continue
        elif "-#-" in line:
            curr_player = line.split("-#-")[0]
            if has_companions:
                companions[curr_player] = []
            else:
                decklists[curr_player] = []
        elif line == "Companions:":
            has_companions = True
        else:
            if has_companions:
                companions[curr_player].append(line)
            else:
                decklists[curr_player].append(line)
    return {
        "color_map" : color_map,
        "color_order" : color_order,
        "companions": companions,
        "date": "",
        "decklists": decklists,
        "draft": "",
        "notes": "",
        "num_players": 4,
        "pack_size": 15,
        "patch": "",
        "player_order": player_order,
        "results": {},
        "rounds": 3,
        "draft_type" : "Pack"
    }
    
def xlsx_to_dict() -> dict:
    """"""
    for f in xlsx_path.iterdir():
        if f.suffix != ".xlsx":
            continue
        try:
            draft_number = int(f.stem)
        except TypeError:
            print(f"{f.stem} is an invalid draft number.")
        
if __name__ == "__main__":
    record_xlsx_draft()